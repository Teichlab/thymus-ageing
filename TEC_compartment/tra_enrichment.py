"""Tissue-restricted antigen (TRA) representation among age-DE genes.

Given the DESeq2 age-DE tables (``{cell_type: df}`` pickles, one per contrast) and the GTEx
tissue-specificity table (``tissue_tau`` / ``max_expr`` per gene), this module answers two
questions per cell type:

1. what fraction of the genes that go DOWN / do not change / go UP with age are TRAs
   (``plot_tra_stacked``), and
2. are TRAs over-represented among the down-regulated genes (``fisher_tra``), and does that
   survive matching on expression level (``cmh_tra``).

Three things about the input data drive the defaults here:

* **The GTEx table has duplicated gene symbols** (199 symbols, up to 8 rows each, with genuinely
  different tau values). A plain ``pd.merge`` on ``gene_name`` therefore duplicates DE rows.
  ``load_gtex_stats`` collapses them (default: keep the most-expressed row per symbol).
* **tau is undefined for genes that are not expressed anywhere.** A gene detected in one tissue at
  0.15 log1p-TPM gets ``tissue_tau == 1.0``, so a bare ``tau >= 0.85`` rule labels thousands of
  noise genes as TRAs. The universe is therefore restricted to genes with
  ``max_expr >= min_max_expr`` (default log1p(1) i.e. >= 1 TPM in the best tissue) *before*
  splitting into TRA / non-TRA.
* **~60% of genes in a DESeq2 table have ``padj == NaN``** (independent filtering). Those genes
  were never tested, so counting them as "does not change" inflates the background. The universe
  defaults to ``padj.notna()``.

Because TRAs are on average much more lowly expressed than non-TRAs (mean log10 baseMean ~0.5 vs
~1.3), while significant genes skew high-expressed, a plain Fisher test is confounded by detection
power in *both* directions. ``cmh_tra`` repeats the test stratified by baseMean decile
(Cochran-Mantel-Haenszel) and is the number to quote.

Main entry points
-----------------
``flatten_de_dict``    - ``{cell_type: df}`` -> long DE frame with a ``change`` column
``load_gtex_stats``    - GTEx tau table, de-duplicated
``build_tra_table``    - join the two, define the universe and ``isTRA``
``tra_counts``         - counts / fractions per cell type x change class
``fisher_tra``         - per-cell-type Fisher exact test + odds ratio with CI, BH-FDR
``cmh_tra``            - the same test stratified by expression level (the control)
``plot_tra_stacked``   - the stacked bar figure
``plot_tra_fraction_bars`` - the TRA share alone, with CIs (more readable than the stack)
``plot_or_forest``     - odds ratios with CI (the direct read-out of the Fisher test)
``plot_tra_running_fraction`` - threshold-free TRA fraction along the DE ranking
``plot_tau_sensitivity``      - TRA fraction per change class vs the tau cut-off
``plot_tra_tissue_composition`` - which tissues the lost TRAs belong to
``down_tra_genes``     - the leading-edge gene table
"""

import sys
from typing import Dict, Iterable, List, Optional, Sequence, Tuple, Union

import numpy as np
import pandas as pd

import matplotlib.pyplot as plt
import seaborn as sns
from matplotlib.lines import Line2D
from matplotlib.patches import Patch

from scipy.stats import fisher_exact, norm
from statsmodels.stats.multitest import multipletests

sys.path.insert(2, '/nfs/team205/vk8/projects/thymus_ageing_atlas/General_analysis/scripts')
from plotting.utils import calc_figsize
from plotting.colors import thyAgeing_colors, thyAgeing_greys


GENERAL_DATA_PATH = '/nfs/team205/vk8/projects/thymus_ageing_atlas/General_analysis/data'
GTEX_STATS_PATH = f'{GENERAL_DATA_PATH}/tra_analysis/GTEX_expression_stats_by_broad_tissue2.csv'

DEFAULT_TAU = 0.85                        # standard TRA cut-off (Yanai et al. tau)
DEFAULT_MIN_MAX_EXPR = float(np.log1p(1))  # >= 1 TPM in the best tissue, ~0.693 on the log1p scale

CHANGE_ORDER = ('sign_dw', 'non_sign', 'sign_up')
CHANGE_LABELS = {'sign_dw': 'Down', 'non_sign': 'No change', 'sign_up': 'Up'}
CHANGE_COLORS = {'sign_dw': thyAgeing_colors['teal'],
                 'non_sign': thyAgeing_greys['grey3'],
                 'sign_up': thyAgeing_colors['orange']}

# TRA / non-TRA fill colours for the stacked bars
TRA_COLORS = {True: thyAgeing_colors['purple'], False: thyAgeing_greys['grey2']}
TRA_LABELS = {True: 'TRA', False: 'non-TRA'}

# GTEx tissues get the ten adopted atlas colours (thyAgeing_colors.csv), assigned in order of how
# many down-regulated TRAs they contribute, with everything past the tenth folded into grey.
TISSUE_PALETTE_ORDER = ['blue', 'orange', 'teal', 'magenta', 'yellow',
                        'purple', 'grass', 'pink', 'mint', 'forest']
OTHER_COLOR = thyAgeing_greys['grey2']


def tissue_palette(tissues: Sequence[str]) -> Dict[str, str]:
    """Map an ordered list of tissue names onto the atlas palette; 'other' is grey."""

    cols = [thyAgeing_colors[c] for c in TISSUE_PALETTE_ORDER]
    pal = {}
    i = 0
    for t in tissues:
        if t == 'other':
            pal[t] = OTHER_COLOR
        else:
            pal[t] = cols[i % len(cols)]
            i += 1
    return pal


_COMPARISONS = {
    # name: (focal change classes, reference change classes)
    'dw_vs_ns':   (('sign_dw',), ('non_sign',)),
    'up_vs_ns':   (('sign_up',), ('non_sign',)),
    'dw_vs_rest': (('sign_dw',), ('non_sign', 'sign_up')),
    'up_vs_rest': (('sign_up',), ('non_sign', 'sign_dw')),
    'dw_vs_up':   (('sign_dw',), ('sign_up',)),
}


# ---------------------------------------------------------------------------------------------
# data preparation
# ---------------------------------------------------------------------------------------------

def flatten_de_dict(de_dict: Dict[str, pd.DataFrame],
                    contrast: str,
                    col_cell_type: str = 'taa_l3',
                    alpha: float = 0.05) -> pd.DataFrame:
    """``{cell_type: DESeq2 result df}`` -> long frame with ``cell_type``/``gene_name``/``change``.

    Same output as the ``_flatten`` helper used in the DE notebooks, kept here so the module can be
    used without it. ``change`` is ``sign_up`` / ``sign_dw`` / ``non_sign`` at ``padj < alpha``
    (genes with ``padj == NaN`` fall in ``non_sign`` and are dropped later by the universe filter).
    """

    df = (pd.concat(list(de_dict.values()), keys=list(de_dict.keys()), names=['cell_type'])
            .reset_index().rename(columns={'level_1': 'gene_name'}))
    df['contrast'] = contrast
    df['col_cell_type'] = col_cell_type
    df['change'] = np.where((df['padj'] < alpha) & (df['log2FoldChange'] > 0), 'sign_up',
                     np.where((df['padj'] < alpha) & (df['log2FoldChange'] < 0), 'sign_dw',
                              'non_sign'))
    return df


def load_gtex_stats(path: str = GTEX_STATS_PATH,
                    dedup: str = 'max_expr') -> pd.DataFrame:
    """Load the GTEx per-gene tissue-specificity table, one row per gene symbol.

    ``dedup`` controls how duplicated symbols are collapsed:
      ``'max_expr'`` keep the row with the highest ``max_expr`` (default - the best-supported copy),
      ``'max_tau'``  keep the row with the highest ``tissue_tau`` (permissive),
      ``'mean'``     average ``tissue_tau``/``max_expr`` (drops ``top1_tissue``),
      ``'drop'``     drop every duplicated symbol.
    """

    stats = pd.read_csv(path, index_col=0)
    stats.index.name = 'gene_name'
    n_dup = int(stats.index.duplicated().sum())

    if dedup == 'max_expr':
        stats = stats.sort_values('max_expr', ascending=False)
        stats = stats[~stats.index.duplicated(keep='first')]
    elif dedup == 'max_tau':
        stats = stats.sort_values('tissue_tau', ascending=False)
        stats = stats[~stats.index.duplicated(keep='first')]
    elif dedup == 'mean':
        stats = stats.groupby(level=0)[['tissue_tau', 'max_expr']].mean()
        stats['top1_tissue'] = np.nan
    elif dedup == 'drop':
        stats = stats[~stats.index.duplicated(keep=False)]
    else:
        raise ValueError(f"dedup must be one of 'max_expr', 'max_tau', 'mean', 'drop'; got {dedup}")

    print(f'GTEx stats: {len(stats)} unique gene symbols '
          f'({n_dup} duplicate rows collapsed with dedup="{dedup}")')
    return stats.sort_index()


# per-gene GENCODE biotype for the atlas var_names, joined through the Ensembl gene_id
# (TEC_compartment/notebooks/notebooks_revisions/_builders/build_gene_biotypes.py)
GENE_BIOTYPES_PATH = f'{GENERAL_DATA_PATH}/tra_analysis/gene_biotypes/atlas_v6_gene_biotypes.csv'


def load_gene_biotypes(path: str = GENE_BIOTYPES_PATH) -> pd.DataFrame:
    """Biotype table for the atlas genes: ``gene_name`` (var_name), ``gene_id``, ``gene_type``,
    ``annotation_source`` and a boolean ``protein_coding``."""

    return pd.read_csv(path)


def load_protein_coding_genes(path: str = GENE_BIOTYPES_PATH, verbose: bool = True) -> set:
    """Atlas var_names whose Ensembl gene record is ``gene_type == 'protein_coding'``.

    Pass the result as ``gene_universe`` to ``build_tra_table`` / ``build_set_table``.
    """

    bt = load_gene_biotypes(path)
    genes = set(bt.loc[bt['protein_coding'].astype(bool), 'gene_name'])
    if verbose:
        src = ', '.join(f'{k}: {v:,}' for k, v in bt['annotation_source'].value_counts().items())
        print(f'protein-coding genes: {len(genes):,} of {len(bt):,} atlas genes ({src})')
    return genes


def build_tra_table(de: Union[pd.DataFrame, Dict[str, pd.DataFrame]],
                    gtex: Optional[pd.DataFrame] = None,
                    tau: float = DEFAULT_TAU,
                    min_max_expr: float = DEFAULT_MIN_MAX_EXPR,
                    require_tested: bool = True,
                    gene_universe: Optional[Iterable[str]] = None,
                    cell_types: Optional[Sequence[str]] = None,
                    contrast: Optional[str] = None,
                    verbose: bool = True,
                    **flatten_kwargs) -> pd.DataFrame:
    """Join the DE table to the GTEx stats, restrict to the testable universe and flag TRAs.

    Parameters
    ----------
    de
        Long DE frame (``cell_type``, ``gene_name``, ``baseMean``, ``log2FoldChange``, ``padj``,
        ``change``) or a ``{cell_type: df}`` dict, which is flattened with ``flatten_de_dict``
        (``contrast`` and any ``flatten_kwargs`` are passed through).
    tau
        ``tissue_tau >= tau`` defines a TRA.
    min_max_expr
        Genes whose best-tissue expression is below this (log1p TPM) are dropped from the universe
        entirely - their tau is meaningless. Set to 0 to keep them (not recommended).
    require_tested
        Keep only genes with a non-NaN ``padj``, i.e. genes DESeq2 actually tested in that cell type.
    gene_universe
        Optional set of gene names (e.g. ``load_protein_coding_genes()``). The universe - TRAs *and*
        the non-TRA background - is restricted to these genes, so the tests compare like with like.
        Filtering only the TRA side would leave lncRNAs and pseudogenes in the background.
    cell_types
        Restrict to these cell types (order is preserved for the plots).

    Returns
    -------
    The universe, one row per (cell type, gene), with ``tissue_tau``, ``max_expr``, ``top1_tissue``
    and a boolean ``isTRA``. ``.attrs['tra_params']`` records the thresholds used.
    """

    if isinstance(de, dict):
        de = flatten_de_dict(de, contrast=contrast or 'unnamed', **flatten_kwargs)
    df = de.copy()

    if 'change' not in df.columns:
        raise KeyError("`de` needs a 'change' column - build it with flatten_de_dict()")
    if cell_types is not None:
        df = df[df['cell_type'].isin(list(cell_types))].copy()
        df['cell_type'] = pd.Categorical(df['cell_type'], categories=list(cell_types), ordered=True)

    gtex = load_gtex_stats() if gtex is None else gtex
    df = df.drop(columns=[c for c in gtex.columns if c in df.columns], errors="ignore")
    keep_cols = [c for c in ['tissue_tau', 'max_expr', 'top1_tissue', 'top5_tissues']
                 if c in gtex.columns]
    n_in = len(df)
    df = df.merge(gtex[keep_cols], left_on='gene_name', right_index=True, how='left')
    if len(df) != n_in:
        raise RuntimeError(f'join duplicated rows ({n_in} -> {len(df)}); pass a de-duplicated '
                           '`gtex` from load_gtex_stats()')

    n_all = len(df)
    in_gtex = df['tissue_tau'].notna()
    expressed = in_gtex & (df['max_expr'] >= min_max_expr)
    tested = df['padj'].notna() if require_tested else pd.Series(True, index=df.index)
    if gene_universe is not None:
        gene_universe = set(gene_universe)
        in_universe = df['gene_name'].isin(gene_universe)
    else:
        in_universe = pd.Series(True, index=df.index)

    out = df[expressed & tested & in_universe].copy()
    out['isTRA'] = out['tissue_tau'] >= tau
    out['change'] = pd.Categorical(out['change'], categories=list(CHANGE_ORDER), ordered=True)
    out.attrs['tra_params'] = dict(tau=tau, min_max_expr=min_max_expr,
                                   require_tested=require_tested,
                                   gene_universe=None if gene_universe is None else len(gene_universe))

    if verbose:
        lines = [f'universe: {len(out):,} of {n_all:,} (cell type, gene) rows kept',
                 f'  - not in GTEx          : {int((~in_gtex).sum()):,}',
                 f'  - max_expr < {min_max_expr:.3f}     : {int((in_gtex & ~expressed).sum()):,}',
                 f'  - padj is NaN          : {int((expressed & ~tested).sum()):,}']
        if gene_universe is not None:
            lines.append(f'  - outside gene universe: '
                         f'{int((expressed & tested & ~in_universe).sum()):,}')
        lines.append(f'  TRAs (tau >= {tau}): {int(out["isTRA"].sum()):,} '
                     f'({100 * out["isTRA"].mean():.1f}% of the universe; '
                     f'{out.loc[out["isTRA"], "gene_name"].nunique():,} distinct genes)')
        print('\n'.join(lines))
    return out


# --------------------------------------------------------------------------------------------
# regulator target sets (mouse AIRE / FEZF2, mapped to human) - see aire_fezf2/raw/SOURCES.md
# --------------------------------------------------------------------------------------------

AIRE_FEZF2_PATH = f'{GENERAL_DATA_PATH}/tra_analysis/aire_fezf2'

REGULATOR_SETS = {
    'AIRE': dict(genes='aire_dependent_human.csv', background='aire_background_human.csv',
                 source='Sansom et al. 2014 Genome Res, Aire+ vs Aire-KO mTEC (>2 fold, 5% FDR)'),
    'FEZF2': dict(genes='fezf2_dependent_human.csv', background='fezf2_background_human.csv',
                  source='Takaba et al. 2015 Cell, re-derived from GSE69105 (FC>2, p<0.05)'),
}


def load_regulator_set(name: str, path: str = AIRE_FEZF2_PATH,
                       verbose: bool = True) -> Tuple[set, set]:
    """Human orthologs of a mouse regulator's target set, plus the assayed background.

    ``name`` is 'AIRE' or 'FEZF2'. Returns ``(targets, background)`` as sets of human symbols.
    The background matters: a gene counts as a non-target only if its mouse ortholog was actually
    measured in that mouse experiment, otherwise "not a target" silently absorbs every gene without
    a 1:1 ortholog and every gene the array or the RNA-seq never quantified.
    """

    if name not in REGULATOR_SETS:
        raise ValueError(f'name must be one of {list(REGULATOR_SETS)}; got {name}')
    spec = REGULATOR_SETS[name]
    targets = set(pd.read_csv(f'{path}/{spec["genes"]}')['human_symbol'])
    background = set(pd.read_csv(f'{path}/{spec["background"]}')['human_symbol'])
    if verbose:
        print(f'{name}: {len(targets):,} target genes, {len(background):,} in the assayed '
              f'background\n  source: {spec["source"]}')
    return targets, background


def build_set_table(de: Union[pd.DataFrame, Dict[str, pd.DataFrame]],
                    gene_set: Union[set, Sequence[str]],
                    background: Optional[Union[set, Sequence[str]]] = None,
                    gtex: Optional[pd.DataFrame] = None,
                    cell_types: Optional[Sequence[str]] = None,
                    require_tested: bool = True,
                    gene_universe: Optional[Iterable[str]] = None,
                    contrast: Optional[str] = None,
                    set_name: str = 'target',
                    verbose: bool = True,
                    **flatten_kwargs) -> pd.DataFrame:
    """``build_tra_table`` for an arbitrary gene set instead of the GTEx tau rule.

    Produces the same schema - a boolean ``isTRA`` column plus ``change``, ``baseMean`` and
    ``cell_type`` - so every test and plot in this module works unchanged on the result.

    Parameters
    ----------
    gene_set
        Human symbols in the set (e.g. AIRE-dependent genes).
    background
        Human symbols that *could* have been called. The universe is restricted to these, so
        ``isTRA == False`` means "assayed and not called" rather than "never looked at". Pass None
        to use every tested gene as background (only sensible if the set is genome-wide by
        construction).
    gene_universe
        Optional further restriction of the universe, e.g. ``load_protein_coding_genes()``.
    gtex
        Optional GTEx stats table; when given, ``tissue_tau`` / ``top1_tissue`` are joined on so the
        tissue-of-origin plots and a TRA cross-tabulation still work.
    """

    if isinstance(de, dict):
        de = flatten_de_dict(de, contrast=contrast or 'unnamed', **flatten_kwargs)
    df = de.copy()

    if 'change' not in df.columns:
        raise KeyError("`de` needs a 'change' column - build it with flatten_de_dict()")
    if cell_types is not None:
        df = df[df['cell_type'].isin(list(cell_types))].copy()
        df['cell_type'] = pd.Categorical(df['cell_type'], categories=list(cell_types), ordered=True)

    if gtex is not None:
        keep_cols = [c for c in ['tissue_tau', 'max_expr', 'top1_tissue', 'top5_tissues']
                     if c in gtex.columns]
        df = df.drop(columns=[c for c in keep_cols if c in df.columns], errors='ignore')
        n_in = len(df)
        df = df.merge(gtex[keep_cols], left_on='gene_name', right_index=True, how='left')
        if len(df) != n_in:
            raise RuntimeError(f'join duplicated rows ({n_in} -> {len(df)}); pass a de-duplicated '
                               '`gtex` from load_gtex_stats()')

    gene_set = set(gene_set)
    n_all = len(df)
    tested = df['padj'].notna() if require_tested else pd.Series(True, index=df.index)
    if background is not None:
        in_bg = df['gene_name'].isin(set(background))
    else:
        in_bg = pd.Series(True, index=df.index)
    if gene_universe is not None:
        in_gu = df['gene_name'].isin(set(gene_universe))
    else:
        in_gu = pd.Series(True, index=df.index)

    out = df[tested & in_bg & in_gu].copy()
    out['isTRA'] = out['gene_name'].isin(gene_set)
    out['change'] = pd.Categorical(out['change'], categories=list(CHANGE_ORDER), ordered=True)
    out.attrs['set_params'] = dict(set_name=set_name, n_set=len(gene_set),
                                   n_background=None if background is None else len(set(background)),
                                   require_tested=require_tested,
                                   gene_universe=None if gene_universe is None else len(set(gene_universe)))

    if verbose:
        n_genes = out['gene_name'].nunique()
        print(f'{set_name} universe: {len(out):,} of {n_all:,} (cell type, gene) rows kept\n'
              f'  - padj is NaN               : {int((~tested).sum()):,}\n'
              f'  - outside assayed background: {int((tested & ~in_bg).sum()):,}\n'
              f'  - outside gene universe     : {int((tested & in_bg & ~in_gu).sum()):,}\n'
              f'  {n_genes:,} distinct genes; {out.loc[:, "isTRA"].sum():,} rows are {set_name}s '
              f'({100 * out["isTRA"].mean():.1f}% of the universe, '
              f'{out.loc[out["isTRA"], "gene_name"].nunique():,} distinct genes)')
    return out


def tra_counts(df: pd.DataFrame,
               cell_type_col: str = 'cell_type') -> pd.DataFrame:
    """Counts and TRA fraction per cell type x change class (tidy, one row per combination)."""

    grp = (df.groupby([cell_type_col, 'change'], observed=False)['isTRA']
             .agg(n_genes='size', n_tra='sum').reset_index())
    grp['n_non_tra'] = grp['n_genes'] - grp['n_tra']
    grp['frac_tra'] = grp['n_tra'] / grp['n_genes'].replace(0, np.nan)
    # binomial (Wilson) 95% CI on the TRA fraction, for the error bars
    lo, hi = _wilson_ci(grp['n_tra'].to_numpy(), grp['n_genes'].to_numpy())
    grp['frac_tra_lo'], grp['frac_tra_hi'] = lo, hi
    return grp


def _wilson_ci(k: np.ndarray, n: np.ndarray, z: float = 1.96) -> Tuple[np.ndarray, np.ndarray]:
    """Wilson score interval for a binomial proportion (stable at k = 0 and k = n)."""

    k, n = np.asarray(k, float), np.asarray(n, float)
    with np.errstate(invalid='ignore', divide='ignore'):
        p = k / n
        denom = 1 + z ** 2 / n
        centre = (p + z ** 2 / (2 * n)) / denom
        half = z * np.sqrt(p * (1 - p) / n + z ** 2 / (4 * n ** 2)) / denom
    return np.where(n > 0, centre - half, np.nan), np.where(n > 0, centre + half, np.nan)


# ---------------------------------------------------------------------------------------------
# statistics
# ---------------------------------------------------------------------------------------------

def _split(df_ct: pd.DataFrame, comparison: str) -> Tuple[pd.Series, pd.Series]:
    """Boolean masks for the focal and reference gene sets of a named comparison."""

    if comparison not in _COMPARISONS:
        raise ValueError(f'comparison must be one of {list(_COMPARISONS)}; got {comparison}')
    focal_lv, ref_lv = _COMPARISONS[comparison]
    ch = df_ct['change'].astype(str)
    return ch.isin(focal_lv), ch.isin(ref_lv)


def _or_ci(a: int, b: int, c: int, d: int, z: float = 1.96) -> Tuple[float, float, float]:
    """Sample odds ratio and Woolf 95% CI with a Haldane-Anscombe 0.5 correction."""

    a_, b_, c_, d_ = a + 0.5, b + 0.5, c + 0.5, d + 0.5
    or_s = (a_ * d_) / (b_ * c_)
    se = np.sqrt(1 / a_ + 1 / b_ + 1 / c_ + 1 / d_)
    return or_s, or_s * np.exp(-z * se), or_s * np.exp(z * se)


def fisher_tra(df: pd.DataFrame,
               comparisons: Union[str, Sequence[str]] = ('dw_vs_ns', 'up_vs_ns'),
               alternative: str = 'two-sided',
               cell_type_col: str = 'cell_type',
               fdr_method: str = 'fdr_bh') -> pd.DataFrame:
    """Fisher exact test for TRA over-representation, per cell type.

    The 2x2 table is (focal / reference gene set) x (TRA / non-TRA); e.g. for ``'dw_vs_ns'`` it asks
    whether the down-with-age genes contain more TRAs than the genes that do not change.

    Returns one row per (cell type, comparison) with the table counts, the TRA fraction in each set,
    the odds ratio with a Haldane-corrected Woolf 95% CI (``odds_ratio_sample``, ``ci_low``,
    ``ci_high``), the exact ``pvalue`` and BH-FDR ``padj`` computed across cell types within each
    comparison.
    """

    comparisons = [comparisons] if isinstance(comparisons, str) else list(comparisons)
    rows = []
    for ct, sub in df.groupby(cell_type_col, observed=True):
        for comp in comparisons:
            focal, ref = _split(sub, comp)
            a = int((focal & sub['isTRA']).sum())       # TRA in focal set
            b = int((focal & ~sub['isTRA']).sum())      # non-TRA in focal set
            c = int((ref & sub['isTRA']).sum())         # TRA in reference set
            d = int((ref & ~sub['isTRA']).sum())        # non-TRA in reference set
            _, p = fisher_exact([[a, b], [c, d]], alternative=alternative)
            or_s, lo, hi = _or_ci(a, b, c, d)
            rows.append(dict(cell_type=ct, comparison=comp,
                             n_focal=a + b, n_ref=c + d, tra_focal=a, tra_ref=c,
                             frac_tra_focal=a / (a + b) if a + b else np.nan,
                             frac_tra_ref=c / (c + d) if c + d else np.nan,
                             odds_ratio_sample=or_s,
                             ci_low=lo, ci_high=hi, pvalue=p))

    res = pd.DataFrame(rows)
    res['padj'] = np.nan
    for comp, idx in res.groupby('comparison').groups.items():
        ok = res.loc[idx, 'pvalue'].notna()
        if ok.any():
            res.loc[idx[ok.to_numpy()], 'padj'] = multipletests(
                res.loc[idx[ok.to_numpy()], 'pvalue'], method=fdr_method)[1]
    res['log2_odds_ratio'] = np.log2(res['odds_ratio_sample'])
    res['stars'] = res['padj'].map(pstars)
    if isinstance(df[cell_type_col].dtype, pd.CategoricalDtype):
        res['cell_type'] = pd.Categorical(res['cell_type'],
                                          categories=df[cell_type_col].cat.categories, ordered=True)
        res = res.sort_values(['comparison', 'cell_type'])
    return res.reset_index(drop=True)


def cmh_tra(df: pd.DataFrame,
            comparison: str = 'dw_vs_ns',
            stratify_on: str = 'baseMean',
            n_bins: int = 10,
            log_strata: bool = True,
            min_focal: int = 10,
            cell_type_col: str = 'cell_type',
            fdr_method: str = 'fdr_bh') -> pd.DataFrame:
    """Expression-matched version of ``fisher_tra``: Cochran-Mantel-Haenszel over baseMean bins.

    TRAs are markedly lower-expressed than non-TRAs while significant genes skew high-expressed, so
    a marginal Fisher test mixes TRA biology with detection power. Here genes are split into
    ``n_bins`` equal-sized bins of ``stratify_on`` *within each cell type*, the 2x2 table is built
    per bin, and the tables are combined with the Mantel-Haenszel common odds ratio and the CMH
    chi-square (hypergeometric mean/variance summed over strata).

    The CMH statistic is a normal approximation and is not trustworthy on a handful of genes: a
    single down-regulated gene that happens to be a TRA returns z ~ 4.8 where the exact
    probability is ~0.04. Cell types with fewer than ``min_focal`` genes in the focal set are
    therefore blanked (NaN) rather than reported - use ``permutation_tra`` for those.

    Returns per cell type: observed vs expected TRA count in the focal set, the MH common odds
    ratio with a Robins-Breslow-Greenland 95% CI, the CMH ``zscore``/``pvalue`` and BH-FDR ``padj``.
    """

    rows = []
    for ct, sub in df.groupby(cell_type_col, observed=True):
        sub = sub.dropna(subset=[stratify_on]).copy()
        x = np.log10(sub[stratify_on] + 1e-9) if log_strata else sub[stratify_on]
        try:
            sub['_stratum'] = pd.qcut(x, n_bins, labels=False, duplicates='drop')
        except ValueError:                                   # too few distinct values
            sub['_stratum'] = 0
        focal, ref = _split(sub, comparison)

        obs = exp = var = 0.0
        num_or = den_or = 0.0                                # MH common odds ratio
        r_sum = s_sum = pr = pspqr = qs = 0.0                # RBG variance pieces
        n_strata = 0
        for _, g in sub.groupby('_stratum'):
            f, r = focal.loc[g.index], ref.loc[g.index]
            a = float((f & g['isTRA']).sum())
            b = float((f & ~g['isTRA']).sum())
            c = float((r & g['isTRA']).sum())
            d = float((r & ~g['isTRA']).sum())
            n = a + b + c + d
            n_focal, n_tra = a + b, a + c
            if n < 2 or n_focal == 0 or n_focal == n or n_tra == 0 or n_tra == n:
                continue                                     # stratum carries no information
            n_strata += 1
            obs += a
            exp += n_focal * n_tra / n
            var += n_focal * (n - n_focal) * n_tra * (n - n_tra) / (n * n * (n - 1))
            R, S = a * d / n, b * c / n
            num_or += R
            den_or += S
            r_sum, s_sum = r_sum + R, s_sum + S
            P, Q = (a + d) / n, (b + c) / n
            pr += P * R
            pspqr += P * S + Q * R
            qs += Q * S

        z = (obs - exp) / np.sqrt(var) if var > 0 else np.nan
        p = 2 * norm.sf(abs(z)) if np.isfinite(z) else np.nan
        or_mh = num_or / den_or if den_or > 0 else np.nan
        if r_sum > 0 and s_sum > 0:
            se_log = np.sqrt(pr / (2 * r_sum ** 2)
                             + pspqr / (2 * r_sum * s_sum)
                             + qs / (2 * s_sum ** 2))
            lo, hi = or_mh * np.exp(-1.96 * se_log), or_mh * np.exp(1.96 * se_log)
        else:
            lo = hi = np.nan
        rows.append(dict(cell_type=ct, comparison=comparison, n_strata=n_strata,
                         n_focal=int(focal.sum()), tra_focal=int(obs), tra_expected=exp,
                         obs_over_exp=obs / exp if exp > 0 else np.nan,
                         odds_ratio_mh=or_mh, ci_low=lo, ci_high=hi, zscore=z, pvalue=p))

    res = pd.DataFrame(rows)
    small = res['n_focal'] < min_focal
    if small.any():
        blanked = sorted(res.loc[small, 'cell_type'].astype(str))
        print(f'cmh_tra: {int(small.sum())} cell type(s) with < {min_focal} genes in the focal set '
              f'blanked - the normal approximation is invalid there: {blanked}')
        res.loc[small, ['odds_ratio_mh', 'ci_low', 'ci_high', 'zscore', 'pvalue']] = np.nan
    ok = res['pvalue'].notna()
    res['padj'] = np.nan
    if ok.any():
        res.loc[ok, 'padj'] = multipletests(res.loc[ok, 'pvalue'], method=fdr_method)[1]
    res['log2_odds_ratio'] = np.log2(res['odds_ratio_mh'])
    res['stars'] = res['padj'].map(pstars)
    if isinstance(df[cell_type_col].dtype, pd.CategoricalDtype):
        res['cell_type'] = pd.Categorical(res['cell_type'],
                                          categories=df[cell_type_col].cat.categories, ordered=True)
        res = res.sort_values('cell_type')
    return res.reset_index(drop=True)


def permutation_tra(df: pd.DataFrame,
                    comparison: str = 'dw_vs_ns',
                    n_perm: int = 2000,
                    stratify_on: Optional[str] = 'baseMean',
                    n_bins: int = 10,
                    cell_type_col: str = 'cell_type',
                    seed: int = 0) -> pd.DataFrame:
    """Empirical null for the TRA count in the focal set, shuffling the TRA label.

    With ``stratify_on`` the label is permuted *within* expression bins, which is the sampling
    version of ``cmh_tra`` and makes no normal approximation - useful when the focal set is small
    (e.g. mTECII, ~50 down genes). Returns observed/expected counts and a two-sided empirical p.
    """

    rng = np.random.default_rng(seed)
    rows = []
    for ct, sub in df.groupby(cell_type_col, observed=True):
        sub = sub.copy()
        focal, _ = _split(sub, comparison)
        keep = focal | _split(sub, comparison)[1]
        sub, focal = sub[keep], focal[keep]
        if stratify_on is not None:
            x = np.log10(sub[stratify_on] + 1e-9)
            try:
                strata = pd.qcut(x, n_bins, labels=False, duplicates='drop').to_numpy()
            except ValueError:
                strata = np.zeros(len(sub), int)
        else:
            strata = np.zeros(len(sub), int)

        is_tra = sub['isTRA'].to_numpy()
        f = focal.to_numpy()
        obs = int((f & is_tra).sum())
        null = np.zeros(n_perm)
        for s in np.unique(strata):
            m = strata == s
            tra_s, f_s = is_tra[m], f[m]
            if f_s.sum() == 0 or tra_s.sum() == 0:
                continue
            perm = np.array([rng.permutation(tra_s) for _ in range(n_perm)])
            null += (perm[:, f_s]).sum(axis=1)
        exp = null.mean()
        # two-sided empirical p, +1 correction
        p = (1 + min((null >= obs).sum(), (null <= obs).sum()) * 2) / (n_perm + 1)
        rows.append(dict(cell_type=ct, comparison=comparison, tra_focal=obs,
                         null_mean=exp, null_sd=null.std(),
                         obs_over_exp=obs / exp if exp > 0 else np.nan,
                         zscore=(obs - exp) / null.std() if null.std() > 0 else np.nan,
                         pvalue=min(p, 1.0)))
    res = pd.DataFrame(rows)
    res['padj'] = multipletests(res['pvalue'], method='fdr_bh')[1]
    res['stars'] = res['padj'].map(pstars)
    return res


def pstars(p: float) -> str:
    """*** < 0.001, ** < 0.01, * < 0.05, 'ns' otherwise."""

    if p is None or not np.isfinite(p):
        return ''
    return '***' if p < 1e-3 else '**' if p < 1e-2 else '*' if p < 0.05 else 'ns'


def tra_summary(df: pd.DataFrame,
                comparison: str = 'dw_vs_ns',
                **cmh_kwargs) -> pd.DataFrame:
    """Fisher and expression-matched CMH results side by side, one row per cell type."""

    fis = fisher_tra(df, comparisons=comparison)
    cmh = cmh_tra(df, comparison=comparison, **cmh_kwargs)
    out = fis[['cell_type', 'n_focal', 'n_ref', 'frac_tra_focal', 'frac_tra_ref',
               'odds_ratio_sample', 'ci_low', 'ci_high', 'pvalue', 'padj']].merge(
        cmh[['cell_type', 'odds_ratio_mh', 'ci_low', 'ci_high', 'tra_focal', 'tra_expected',
             'pvalue', 'padj']],
        on='cell_type', how='left', suffixes=('_fisher', '_matched'))
    out = out.rename(columns={'odds_ratio_sample': 'OR_fisher', 'odds_ratio_mh': 'OR_matched',
                              'tra_focal': 'obs_tra', 'tra_expected': 'exp_tra_matched'})
    out['OR_fisher_ci'] = [f'{lo:.2f}-{hi:.2f}'
                           for lo, hi in zip(out['ci_low_fisher'], out['ci_high_fisher'])]
    out['OR_matched_ci'] = [f'{lo:.2f}-{hi:.2f}'
                            for lo, hi in zip(out['ci_low_matched'], out['ci_high_matched'])]
    out['stars_matched'] = out['padj_matched'].map(pstars)
    out['stars_fisher'] = out['padj_fisher'].map(pstars)
    cols = ['cell_type', 'n_focal', 'n_ref', 'frac_tra_focal', 'frac_tra_ref',
            'OR_fisher', 'OR_fisher_ci', 'pvalue_fisher', 'padj_fisher', 'stars_fisher',
            'obs_tra', 'exp_tra_matched', 'OR_matched', 'OR_matched_ci',
            'pvalue_matched', 'padj_matched', 'stars_matched']
    return out[cols]


def down_tra_genes(df: pd.DataFrame,
                   cell_types: Optional[Sequence[str]] = None,
                   change: str = 'sign_dw',
                   sort_by: str = 'log2FoldChange') -> pd.DataFrame:
    """The leading edge: TRAs in a given change class, with their GTEx tissue of origin."""

    sub = df[(df['change'].astype(str) == change) & df['isTRA']]
    if cell_types is not None:
        sub = sub[sub['cell_type'].isin(list(cell_types))]
    cols = [c for c in ['cell_type', 'gene_name', 'baseMean', 'log2FoldChange', 'padj',
                        'tissue_tau', 'max_expr', 'top1_tissue'] if c in sub.columns]
    return sub[cols].sort_values(['cell_type', sort_by]).reset_index(drop=True)


# --------------------------------------------------------------------------------------------
# supplementary tables (manuscript)
# --------------------------------------------------------------------------------------------

SUPP_DIRECTION_LABELS = {'sign_dw': 'Down with age', 'sign_up': 'Up with age'}

# (source column, manuscript header, legend). '{contrast}' is filled in at build time.
SUPP_GENE_COLUMNS = [
    ('cell_type', 'Cell type', 'TEC population (annotation level taa_l4).'),
    ('gene_name', 'Gene symbol', 'Gene symbol in the atlas gene annotation.'),
    ('gene_id', 'Ensembl gene ID', 'Ensembl gene identifier of the atlas gene record (version stripped).'),
    ('gene_type', 'Gene biotype', 'GENCODE gene_type of the gene record.'),
    ('direction', 'Direction with age', 'Down with age: lower in the older group; Up with age: higher in '
                                        'the older group (sign of the log2 fold change).'),
    ('log2FoldChange', 'log2 fold change ({contrast})', 'DESeq2 log2 fold change, {contrast}.'),
    ('lfcSE', 'log2 fold change SE', 'Standard error of the log2 fold change (DESeq2).'),
    ('stat', 'Wald statistic', 'DESeq2 Wald statistic.'),
    ('pvalue', 'P value', 'DESeq2 Wald test P value.'),
    ('padj', 'Adjusted P value (BH)', 'Benjamini-Hochberg adjusted P value (DESeq2); significant at < 0.05.'),
    ('baseMean', 'baseMean', 'Mean of normalised counts across samples (DESeq2).'),
    ('tissue_tau', 'GTEx tissue specificity (tau)', 'Tissue specificity index tau across GTEx broad '
                                                    'tissues (0 = ubiquitous, 1 = one tissue only).'),
    ('max_expr_tpm', 'GTEx maximum expression (TPM)', 'Expression in the highest-expressing GTEx broad '
                                                      'tissue (TPM).'),
    ('top1_tissue', 'GTEx top tissue', 'GTEx broad tissue with the highest expression.'),
    ('top5', 'GTEx top 5 tissues', 'The five highest-expressing GTEx broad tissues, in decreasing order.'),
]


def _clean_tissue_list(x) -> str:
    """GTEx 'top5_tissues' is stored as the repr of a Python list; render it as 'A; B; C'."""

    import ast
    if isinstance(x, str) and x.startswith('['):
        try:
            return '; '.join(ast.literal_eval(x))
        except (ValueError, SyntaxError):
            return x
    return '' if x is None or (isinstance(x, float) and np.isnan(x)) else str(x)


def supplementary_tra_table(df: pd.DataFrame,
                            cell_types: Sequence[str],
                            changes: Sequence[str] = ('sign_dw', 'sign_up'),
                            biotypes: Optional[pd.DataFrame] = None,
                            regulator_sets: Optional[Dict[str, Tuple[set, set]]] = None,
                            contrast_label: str = 'adult vs infant') -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Manuscript table: one row per TRA significantly changed with age, per cell type.

    Parameters
    ----------
    df
        A ``build_tra_table`` result (it must still carry the DESeq2 columns).
    biotypes
        ``load_gene_biotypes()``; adds the Ensembl ID and GENCODE biotype.
    regulator_sets
        ``{'AIRE': (targets, background), ...}`` from ``load_regulator_set``; adds one column per
        regulator with 'yes' / 'no' / 'not assayed' (no 1:1 ortholog measured in that mouse dataset).

    Returns
    -------
    ``(table, legend)`` - the table with manuscript headers, sorted by cell type, direction and
    decreasing absolute fold change, and a two-column legend describing every column.
    """

    sub = df[df['isTRA'] & df['change'].astype(str).isin(list(changes))].copy()
    sub = sub[sub['cell_type'].astype(str).isin(list(cell_types))]
    sub['cell_type'] = pd.Categorical(sub['cell_type'].astype(str), categories=list(cell_types),
                                      ordered=True)
    sub['direction'] = sub['change'].astype(str).map(SUPP_DIRECTION_LABELS)
    if 'max_expr' in sub.columns:
        sub['max_expr_tpm'] = np.expm1(sub['max_expr'])
    if 'top5_tissues' in sub.columns:
        sub['top5'] = sub['top5_tissues'].map(_clean_tissue_list)
    if biotypes is not None:
        bt = biotypes.drop_duplicates('gene_name').set_index('gene_name')
        sub['gene_id'] = sub['gene_name'].map(bt['gene_id'])
        sub['gene_type'] = sub['gene_name'].map(bt['gene_type'])

    sub['_dir'] = sub['change'].astype(str).map({'sign_dw': 0, 'sign_up': 1})
    sub['_abs'] = -sub['log2FoldChange'].abs()
    sub = sub.sort_values(['cell_type', '_dir', '_abs'])

    cols, headers, legend = [], [], []
    for src, header, desc in SUPP_GENE_COLUMNS:
        if src in sub.columns:
            cols.append(src)
            headers.append(header.format(contrast=contrast_label))
            legend.append((header.format(contrast=contrast_label), desc.format(contrast=contrast_label)))
    table = sub[cols].copy()
    table.columns = headers
    table['Cell type'] = table['Cell type'].astype(str)

    for name, (targets, background) in (regulator_sets or {}).items():
        header = f'{name}-dependent in mouse'
        genes = sub['gene_name']
        table[header] = np.where(~genes.isin(set(background)), 'not assayed',
                                 np.where(genes.isin(set(targets)), 'yes', 'no'))
        src = REGULATOR_SETS.get(name, {}).get('source', '')
        legend.append((header, f'Human 1:1 ortholog of a mouse {name}-dependent gene ({src}); '
                               '"not assayed" = no 1:1 ortholog measured in that mouse dataset.'))

    return table.reset_index(drop=True), pd.DataFrame(legend, columns=['Column', 'Description'])


def supplementary_summary_table(df: pd.DataFrame,
                                cell_types: Sequence[str],
                                n_perm: int = 10000,
                                seed: int = 0) -> pd.DataFrame:
    """Manuscript table: per-population gene / TRA counts and the enrichment tests.

    FDRs are Benjamini-Hochberg across the populations passed in ``cell_types``.
    """

    d = df[df['cell_type'].astype(str).isin(list(cell_types))].copy()
    d['cell_type'] = pd.Categorical(d['cell_type'].astype(str), categories=list(cell_types),
                                    ordered=True)
    cnt = tra_counts(d)
    cnt['cell_type'] = cnt['cell_type'].astype(str)
    cnt['change'] = cnt['change'].astype(str)
    ng = cnt.pivot(index='cell_type', columns='change', values='n_genes').reindex(list(cell_types))
    nt = cnt.pivot(index='cell_type', columns='change', values='n_tra').reindex(list(cell_types))

    out = pd.DataFrame(index=list(cell_types))
    out['Genes tested'] = ng.sum(axis=1).astype(int)
    out['TRAs tested'] = nt.sum(axis=1).astype(int)
    out['TRA fraction, all tested genes (%)'] = 100 * out['TRAs tested'] / out['Genes tested']
    for ch, lab in [('sign_dw', 'Down with age'), ('non_sign', 'No change'), ('sign_up', 'Up with age')]:
        out[f'{lab}: genes'] = ng[ch].fillna(0).astype(int)
        out[f'{lab}: TRAs'] = nt[ch].fillna(0).astype(int)
        out[f'{lab}: TRA fraction (%)'] = 100 * nt[ch] / ng[ch].replace(0, np.nan)

    def _by_ct(frame):
        return (frame.assign(cell_type=frame['cell_type'].astype(str))
                     .set_index('cell_type').reindex(list(cell_types)))

    fis = fisher_tra(d, comparisons=('dw_vs_ns', 'up_vs_ns'))
    for comp, lab in [('dw_vs_ns', 'Down vs no change'), ('up_vs_ns', 'Up vs no change')]:
        f = _by_ct(fis[fis['comparison'] == comp])
        out[f'{lab}: odds ratio (Fisher)'] = f['odds_ratio_sample']
        out[f'{lab}: 95% CI'] = [f'{lo:.2f}-{hi:.2f}' if np.isfinite(lo) else ''
                                 for lo, hi in zip(f['ci_low'], f['ci_high'])]
        out[f'{lab}: P value (Fisher)'] = f['pvalue']
        out[f'{lab}: FDR (Fisher)'] = f['padj']
        if comp == 'dw_vs_ns':
            cm = _by_ct(cmh_tra(d, comparison=comp))
            out[f'{lab}: odds ratio (expression-matched)'] = cm['odds_ratio_mh']
            out[f'{lab}: FDR (expression-matched)'] = cm['padj']
            pm = _by_ct(permutation_tra(d, comparison=comp, n_perm=n_perm,
                                        stratify_on='baseMean', seed=seed))
            out[f'{lab}: P value (permutation, expression-matched)'] = pm['pvalue']
    return out.reset_index().rename(columns={'index': 'Cell type'})


def _sheet_name(name: str) -> str:
    bad = '[]:*?/\\'
    return ''.join('-' if c in bad else c for c in str(name))[:31]


def write_supplementary_xlsx(path: str,
                             sheets: Dict[str, pd.DataFrame],
                             readme: Optional[pd.DataFrame] = None) -> str:
    """Write a manuscript supplementary workbook: optional README first, then one sheet per table.

    Headers are bold and wrapped, the header row is frozen, data sheets get an autofilter, column
    widths follow the content, P values / FDRs are shown in scientific notation and other floats
    to three decimals.
    """

    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        if readme is not None:
            readme.to_excel(xw, sheet_name='README', index=False)
        for name, t in sheets.items():
            t.to_excel(xw, sheet_name=_sheet_name(name), index=False)

        for ws in xw.book.worksheets:
            is_readme = ws.title == 'README'
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            ws.freeze_panes = 'A2'
            if not is_readme and ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
                header = str(col[0].value or '')
                sample = [len(str(c.value)) for c in col[1:501] if c.value is not None]
                if is_readme:
                    width = 34 if j == 1 else 110
                else:
                    width = min(max([len(header) // 2 + 1] + sample + [10]) + 2, 50)
                ws.column_dimensions[get_column_letter(j)].width = width
                is_p = ('P value' in header) or ('FDR' in header)
                for c in col[1:]:
                    if is_readme:
                        c.alignment = Alignment(wrap_text=True, vertical='top')
                    elif isinstance(c.value, float):
                        c.number_format = '0.00E+00' if is_p else '0.000'
    print(f'wrote {path}: ' + ', '.join(f'{_sheet_name(k)} ({len(v)} rows)' for k, v in sheets.items()))
    return path


# ---------------------------------------------------------------------------------------------
# plots
# ---------------------------------------------------------------------------------------------

def _cell_type_order(df: pd.DataFrame,
                     cell_types: Optional[Sequence[str]],
                     col: str = 'cell_type') -> List[str]:
    if cell_types is not None:
        return [c for c in cell_types if c in set(df[col].astype(str))]
    if isinstance(df[col].dtype, pd.CategoricalDtype):
        present = set(df[col].astype(str))
        return [c for c in df[col].cat.categories if c in present]
    return sorted(df[col].astype(str).unique())


def _finish(fig, save: Optional[str], show: bool = True):
    if save is not None:
        fig.savefig(save, bbox_inches='tight')
        print(f'saved {save}')
    if show:
        plt.show()
    return fig


def plot_tra_stacked(df: pd.DataFrame,
                     cell_types: Optional[Sequence[str]] = None,
                     mode: str = 'fraction',
                     stats: Optional[pd.DataFrame] = None,
                     set_name: str = 'TRA',
                     annotate_n: bool = True,
                     show_universe_line: bool = True,
                     bar_width: float = 0.85,
                     group_gap: float = 0.9,
                     figsize: Optional[Tuple[float, float]] = None,
                     ax: Optional[plt.Axes] = None,
                     title: Optional[str] = None,
                     save: Optional[str] = None,
                     show: bool = True):
    """Stacked bars of TRA vs non-TRA within each change class, grouped by cell type.

    Three bars per cell type - genes DOWN with age, genes that do not change, genes UP with age -
    each split into the TRA and non-TRA share.

    Parameters
    ----------
    mode
        ``'fraction'`` (default) - each bar sums to 1, so the bars are directly comparable;
        ``'count'`` - stacked absolute gene counts (dominated by the non-significant class).
    stats
        Optional result of ``fisher_tra``/``cmh_tra``/``tra_summary`` for a down-vs-non-significant
        comparison; its ``stars`` are drawn above each cell-type group.
    annotate_n
        Print the number of genes in each bar, inside the bar at the top (outside the bar in
        ``'count'`` mode, where the bars have no common ceiling).
    show_universe_line
        Dashed line at the TRA fraction of the whole tested universe (fraction mode only).
    """

    order = _cell_type_order(df, cell_types)
    cnt = tra_counts(df)
    cnt = cnt[cnt['cell_type'].astype(str).isin(order)]
    cnt['cell_type'] = pd.Categorical(cnt['cell_type'].astype(str), categories=order, ordered=True)

    n_ch = len(CHANGE_ORDER)
    step = n_ch + group_gap
    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=min(183, 22 + 13 * len(order)), height=55, margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    stars = {}
    if stats is not None:
        star_col = 'stars_matched' if 'stars_matched' in stats.columns else 'stars'
        stars = dict(zip(stats['cell_type'].astype(str), stats[star_col].fillna('')))

    xs, xlabels, xcolors = [], [], []
    for i, ct in enumerate(order):
        for j, ch in enumerate(CHANGE_ORDER):
            row = cnt[(cnt['cell_type'].astype(str) == ct) & (cnt['change'].astype(str) == ch)]
            if row.empty:
                continue
            row = row.iloc[0]
            x = i * step + j
            n, k = int(row['n_genes']), int(row['n_tra'])
            if mode == 'fraction':
                bottom, top = (k / n, 1 - k / n) if n else (0, 0)
            elif mode == 'count':
                bottom, top = k, n - k
            else:
                raise ValueError("mode must be 'fraction' or 'count'")
            ax.bar(x, bottom, width=bar_width, color=TRA_COLORS[True],
                   edgecolor='k', linewidth=0.4, zorder=2)
            ax.bar(x, top, bottom=bottom, width=bar_width, color=TRA_COLORS[False],
                   edgecolor='k', linewidth=0.4, zorder=2)
            if annotate_n:
                # inside the bar in fraction mode, so the space above stays free for the stars
                if mode == 'fraction':
                    ax.text(x, 0.985, f'{n:,}', ha='center', va='top', fontsize=4.5,
                            rotation=90, color=thyAgeing_greys['grey4'], zorder=4)
                else:
                    ax.text(x, n * 1.02, f'{n:,}', ha='center', va='bottom', fontsize=4.5,
                            rotation=90, color='k', zorder=4)
            # % TRA inside the bar when there is room
            if mode == 'fraction' and n and k / n > 0.10:
                ax.text(x, k / n / 2, f'{100 * k / n:.0f}', ha='center', va='center',
                        fontsize=5, color='white', zorder=3)
            xs.append(x)
            xlabels.append(CHANGE_LABELS[ch])
            xcolors.append(CHANGE_COLORS[ch])

        # cell type label under the group, and the significance star above it
        ax.text(i * step + (n_ch - 1) / 2, -0.30, ct, ha='right', va='top', rotation=45,
                fontsize=7, transform=ax.get_xaxis_transform())
        if ct in stars and stars[ct] not in ('', 'ns'):
            ax.text(i * step + (n_ch - 1) / 2, 1.02, stars[ct], ha='center', va='bottom',
                    fontsize=8, transform=ax.get_xaxis_transform())

    handles = [Patch(facecolor=TRA_COLORS[True], edgecolor='k', lw=0.4, label=set_name),
               Patch(facecolor=TRA_COLORS[False], edgecolor='k', lw=0.4, label=f'non-{set_name}')]
    if show_universe_line and mode == 'fraction':
        f0 = df['isTRA'].mean()
        ax.axhline(f0, color='k', ls='--', lw=0.5, zorder=4)
        handles.append(Line2D([], [], color='k', ls='--', lw=0.5,
                              label=f'all tested genes\n({100 * f0:.1f}% {set_name})'))

    ax.set_xticks(xs)
    ax.set_xticklabels(xlabels, rotation=90, fontsize=6)
    for tick, col in zip(ax.get_xticklabels(), xcolors):
        tick.set_color(col)
    ax.tick_params(axis='x', length=0, pad=1)
    ax.set_ylabel('Fraction of genes' if mode == 'fraction' else 'Number of genes')
    if mode == 'fraction':
        ax.set_ylim(0, 1)
    ax.set_xlim(-0.5 - (1 - bar_width), (len(order) - 1) * step + n_ch - 0.5 + (1 - bar_width))
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    if title:
        ax.set_title(title, pad=14)

    ax.legend(handles=handles, frameon=False, fontsize=6, loc='upper left',
              bbox_to_anchor=(1.005, 1.0), handlelength=1.2)
    return _finish(fig, save, show)


def plot_tra_fraction_bars(df: pd.DataFrame,
                           cell_types: Optional[Sequence[str]] = None,
                           stats: Optional[pd.DataFrame] = None,
                           set_name: str = 'TRA',
                           ylabel: Optional[str] = None,
                           annotate_n: bool = True,
                           bar_width: float = 0.27,
                           figsize: Optional[Tuple[float, float]] = None,
                           ax: Optional[plt.Axes] = None,
                           title: Optional[str] = None,
                           save: Optional[str] = None,
                           show: bool = True):
    """The TRA share alone, as dodged bars with 95% CIs - the readable version of the stacked bars.

    A 100%-stacked bar spends ~90% of its ink on the non-TRA remainder, which carries no
    information once the TRA share is drawn. Here only the TRA fraction is plotted, one bar per
    change class, with a Wilson 95% interval (so the small classes - mTECII has ~50 down-regulated
    genes - are visibly uncertain) and the whole-universe TRA fraction as a dashed reference.
    """

    order = _cell_type_order(df, cell_types)
    cnt = tra_counts(df)
    cnt = cnt[cnt['cell_type'].astype(str).isin(order)]

    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=min(183, 25 + 15 * len(order)), height=50, margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    stars = {}
    if stats is not None:
        star_col = "stars_matched" if "stars_matched" in stats.columns else "stars"
        stars = dict(zip(stats["cell_type"].astype(str), stats[star_col].fillna("")))

    offs = np.linspace(-bar_width, bar_width, len(CHANGE_ORDER))
    for j, (ch, off) in enumerate(zip(CHANGE_ORDER, offs)):
        sub = (cnt[cnt['change'].astype(str) == ch]
               .set_index(cnt.loc[cnt['change'].astype(str) == ch, 'cell_type'].astype(str))
               .reindex(order))
        x = np.arange(len(order)) + off
        ax.bar(x, sub['frac_tra'], width=bar_width, color=CHANGE_COLORS[ch],
               edgecolor='k', linewidth=0.4, zorder=2, label=CHANGE_LABELS[ch])
        ax.errorbar(x, sub['frac_tra'],
                    yerr=[sub['frac_tra'] - sub['frac_tra_lo'], sub['frac_tra_hi'] - sub['frac_tra']],
                    fmt='none', ecolor='k', elinewidth=0.5, capsize=1, zorder=3)
        if annotate_n:
            for xi, n, hi in zip(x, sub['n_genes'], sub['frac_tra_hi']):
                if np.isfinite(hi):
                    ax.text(xi, hi + 0.005, f'{int(n):,}', ha='center', va='bottom',
                            fontsize=4.5, rotation=90)

    f0 = df["isTRA"].mean()
    ax.axhline(f0, color="k", ls="--", lw=0.5, zorder=1)

    ymax = np.nanmax(cnt['frac_tra_hi'].to_numpy())
    for i, ct in enumerate(order):
        if ct in stars and stars[ct] not in ('', 'ns'):
            ax.text(i, ymax * 1.14, stars[ct], ha='center', va='bottom', fontsize=8)

    ax.set_xticks(range(len(order)))
    ax.set_xticklabels(order, rotation=45, ha='right', fontsize=7)
    ax.set_ylabel(ylabel if ylabel is not None else f'Fraction of genes that are {set_name}s')
    ax.set_ylim(0, ymax * 1.22)
    ax.set_xlim(-0.6, len(order) - 0.4)
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    if title:
        ax.set_title(title)
    handles = [Patch(facecolor=CHANGE_COLORS[c], edgecolor="k", lw=0.4, label=CHANGE_LABELS[c])
               for c in CHANGE_ORDER]
    handles.append(Line2D([], [], color='k', ls='--', lw=0.5,
                          label=f'all tested genes\n({100 * f0:.1f}% {set_name})'))
    ax.legend(handles=handles, frameon=False, fontsize=6, loc="upper left",
              bbox_to_anchor=(1.005, 1.0), handlelength=1.2, title="with age", title_fontsize=6)
    return _finish(fig, save, show)


def plot_or_forest(results: Union[pd.DataFrame, Dict[str, pd.DataFrame]],
                   or_col: Optional[str] = None,
                   cell_types: Optional[Sequence[str]] = None,
                   colors: Optional[Dict[str, str]] = None,
                   xlim: Optional[Tuple[float, float]] = None,
                   annotate_p: bool = True,
                   figsize: Optional[Tuple[float, float]] = None,
                   ax: Optional[plt.Axes] = None,
                   title: Optional[str] = None,
                   xlabel: str = 'Odds ratio (TRA among down-regulated genes)',
                   save: Optional[str] = None,
                   show: bool = True):
    """Odds ratios with 95% CI, one row per cell type - the direct read-out of the Fisher test.

    ``results`` is a ``fisher_tra``/``cmh_tra`` frame, or a ``{label: frame}`` dict to overlay
    several versions of the test (e.g. ``{'raw': fisher, 'expression-matched': cmh}``), which is
    the honest way to show that the enrichment is not a detection-power artefact. Filled markers
    are FDR < 0.05, open markers are not.
    """

    if isinstance(results, pd.DataFrame):
        results = {'': results}
    labels = list(results)
    if colors is None:
        base = [thyAgeing_colors['magenta'], thyAgeing_colors['blue'],
                thyAgeing_colors['grass'], thyAgeing_colors['yellow']]
        colors = {lab: base[i % len(base)] for i, lab in enumerate(labels)}

    first = results[labels[0]]
    order = _cell_type_order(first, cell_types)[::-1]      # top of the plot = first cell type

    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=70, height=max(28, 9 * len(order)), margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    offsets = np.linspace(-0.18, 0.18, len(labels)) if len(labels) > 1 else [0.0]
    for lab, off in zip(labels, offsets):
        res = results[lab].set_index(results[lab]['cell_type'].astype(str))
        col = or_col or ('odds_ratio_mh' if 'odds_ratio_mh' in res.columns else 'odds_ratio_sample')
        for i, ct in enumerate(order):
            if ct not in res.index:
                continue
            r = res.loc[ct]
            sig = np.isfinite(r.get('padj', np.nan)) and r['padj'] < 0.05
            ax.plot([r['ci_low'], r['ci_high']], [i + off] * 2, color=colors[lab], lw=0.8,
                    solid_capstyle='butt', zorder=2)
            ax.plot(r[col], i + off, 'o', ms=3.5, color=colors[lab], zorder=3,
                    markerfacecolor=colors[lab] if sig else 'white', markeredgewidth=0.8)
            if annotate_p and lab == labels[-1]:
                ax.text(1.02, i, r.get('stars', ''), transform=ax.get_yaxis_transform(),
                        va='center', ha='left', fontsize=6)

    ax.axvline(1, color='k', ls='--', lw=0.5, zorder=1)
    ax.set_xscale('log')
    if xlim:
        ax.set_xlim(*xlim)
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(order, fontsize=7)
    ax.set_ylim(-0.6, len(order) - 0.4)
    ax.set_xlabel(xlabel)
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    if title:
        ax.set_title(title)
    if len(labels) > 1:
        ax.legend(handles=[Line2D([], [], color=colors[l], marker='o', ms=3.5, lw=0.8, label=l)
                           for l in labels],
                  frameon=False, fontsize=6, loc='upper left', bbox_to_anchor=(1.05, 1.0))
    return _finish(fig, save, show)


def plot_tra_running_fraction(df: pd.DataFrame,
                              cell_types: Optional[Sequence[str]] = None,
                              rank_by: str = 'stat',
                              window: int = 500,
                              palette: Optional[Dict[str, str]] = None,
                              figsize: Optional[Tuple[float, float]] = None,
                              ax: Optional[plt.Axes] = None,
                              title: Optional[str] = None,
                              save: Optional[str] = None,
                              show: bool = True):
    """Threshold-free view: TRA fraction along the DE ranking, in a sliding window.

    Genes are ranked by ``rank_by`` (default the DESeq2 Wald statistic: most down-regulated on the
    left, most up-regulated on the right) and the local TRA fraction is drawn against the rank
    percentile, with the cell type's overall TRA fraction as a dashed reference. A curve that rises
    at the left edge is the same signal the Fisher test picks up, without the padj < 0.05 cut.
    """

    order = _cell_type_order(df, cell_types)
    if palette is None:
        cols = plt.get_cmap('tab10').colors
        palette = {ct: cols[i % 10] for i, ct in enumerate(order)}

    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=89, height=55, margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    for ct in order:
        sub = df[df['cell_type'].astype(str) == ct].dropna(subset=[rank_by])
        if len(sub) < window * 2:
            continue
        sub = sub.sort_values(rank_by)
        tra = sub['isTRA'].to_numpy(float)
        run = pd.Series(tra).rolling(window, center=True, min_periods=window // 2).mean()
        pct = 100 * np.arange(len(tra)) / (len(tra) - 1)
        ax.plot(pct, run, color=palette[ct], lw=0.9, label=ct)
        ax.axhline(tra.mean(), color=palette[ct], lw=0.4, ls=':', alpha=0.6)

    ax.set_xlabel(f'Gene rank percentile (by {rank_by}: down $\\rightarrow$ up with age)')
    ax.set_ylabel(f'TRA fraction\n(sliding window, {window} genes)')
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    ax.legend(frameon=False, fontsize=6, loc='upper left', bbox_to_anchor=(1.01, 1.0),
              handlelength=1.2)
    if title:
        ax.set_title(title)
    return _finish(fig, save, show)


def tau_sensitivity_table(df: pd.DataFrame,
                          taus: Sequence[float] = (0.70, 0.75, 0.80, 0.85, 0.90, 0.95),
                          comparison: str = 'dw_vs_ns',
                          matched: bool = True) -> pd.DataFrame:
    """Re-run the enrichment test at a range of tau cut-offs; long frame of odds ratios and FDR."""

    out = []
    for tau in taus:
        d = df.copy()
        d['isTRA'] = d['tissue_tau'] >= tau
        res = (cmh_tra(d, comparison=comparison) if matched
               else fisher_tra(d, comparisons=comparison))
        res = res.rename(columns={'odds_ratio_mh': 'odds_ratio',
                                  'odds_ratio_sample': 'odds_ratio'})
        res['tau'] = tau
        res['n_tra'] = int(d['isTRA'].sum())
        out.append(res[['tau', 'cell_type', 'odds_ratio', 'pvalue', 'padj', 'n_tra']])
    return pd.concat(out, ignore_index=True)


def plot_tau_sensitivity(df: pd.DataFrame,
                         taus: Sequence[float] = (0.70, 0.75, 0.80, 0.85, 0.90, 0.95),
                         comparison: str = 'dw_vs_ns',
                         matched: bool = True,
                         cell_types: Optional[Sequence[str]] = None,
                         palette: Optional[Dict[str, str]] = None,
                         table: Optional[pd.DataFrame] = None,
                         figsize: Optional[Tuple[float, float]] = None,
                         ax: Optional[plt.Axes] = None,
                         title: Optional[str] = None,
                         save: Optional[str] = None,
                         show: bool = True):
    """Odds ratio of the enrichment as a function of the tau cut-off - is 0.85 a lucky threshold?

    ``matched=True`` uses the expression-matched (CMH) odds ratio, otherwise the plain Fisher one.
    Filled markers mark FDR < 0.05 at that cut-off. Pass a precomputed ``table`` from
    ``tau_sensitivity_table`` to avoid re-running the tests.
    """

    order = _cell_type_order(df, cell_types)
    if palette is None:
        cols = plt.get_cmap('tab10').colors
        palette = {ct: cols[i % 10] for i, ct in enumerate(order)}
    if table is None:
        table = tau_sensitivity_table(df, taus=taus, comparison=comparison, matched=matched)

    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=89, height=55, margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    for ct in order:
        sub = table[table['cell_type'].astype(str) == ct].sort_values('tau')
        if sub.empty:
            continue
        ax.plot(sub['tau'], sub['odds_ratio'], color=palette[ct], lw=0.8, label=ct, zorder=1)
        sig = sub['padj'] < 0.05
        ax.plot(sub.loc[sig, 'tau'], sub.loc[sig, 'odds_ratio'], 'o', ms=3,
                color=palette[ct], markerfacecolor=palette[ct], markeredgewidth=0.7, zorder=2)
        ax.plot(sub.loc[~sig, 'tau'], sub.loc[~sig, 'odds_ratio'], 'o', ms=3,
                color=palette[ct], markerfacecolor='white', markeredgewidth=0.7, zorder=2)

    ax.axhline(1, color='k', ls='--', lw=0.5)
    ax.axvline(DEFAULT_TAU, color=thyAgeing_greys['grey3'], ls=':', lw=0.5)
    ax.set_yscale('log')
    ax.set_xlabel(r'TRA definition: tissue $\tau \geq$')
    ax.set_ylabel('Odds ratio' + (' (expression-matched)' if matched else ''))
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    ax.legend(frameon=False, fontsize=6, loc='upper left', bbox_to_anchor=(1.01, 1.0),
              handlelength=1.2)
    if title:
        ax.set_title(title)
    return _finish(fig, save, show)


def plot_tra_tissue_composition(df: pd.DataFrame,
                                cell_types: Optional[Sequence[str]] = None,
                                change: str = 'sign_dw',
                                n_tissues: int = 10,
                                normalise: bool = True,
                                figsize: Optional[Tuple[float, float]] = None,
                                ax: Optional[plt.Axes] = None,
                                title: Optional[str] = None,
                                save: Optional[str] = None,
                                show: bool = True):
    """Which peripheral tissues the TRAs in a change class belong to (GTEx ``top1_tissue``)."""

    order = _cell_type_order(df, cell_types)
    sub = df[(df['change'].astype(str) == change) & df['isTRA']].copy()
    sub = sub[sub['cell_type'].astype(str).isin(order)]
    top = sub['top1_tissue'].value_counts().head(n_tissues).index.tolist()
    sub['tissue'] = np.where(sub['top1_tissue'].isin(top), sub['top1_tissue'], 'other')

    tab = (pd.crosstab(sub['cell_type'].astype(str), sub['tissue'])
             .reindex(index=order, columns=top + ['other'], fill_value=0))
    if normalise:
        tab = tab.div(tab.sum(axis=1).replace(0, np.nan), axis=0)

    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=89, height=55, margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    pal = tissue_palette(list(tab.columns))
    bottom = np.zeros(len(tab))
    for i, t in enumerate(tab.columns):
        c = pal[t]
        ax.bar(range(len(tab)), tab[t].to_numpy(), bottom=bottom, width=0.8,
               color=c, edgecolor='k', linewidth=0.3, label=t)
        bottom += np.nan_to_num(tab[t].to_numpy())

    ax.set_xticks(range(len(tab)))
    ax.set_xticklabels(tab.index, rotation=90, fontsize=7)
    ax.set_ylabel('Fraction of TRAs' if normalise else 'Number of TRAs')
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    ax.legend(frameon=False, fontsize=5.5, loc='upper left', bbox_to_anchor=(1.01, 1.0),
              handlelength=1.0, ncol=1)
    ax.set_title(title or f'GTEx tissue of origin, {CHANGE_LABELS.get(change, change)} TRAs')
    return _finish(fig, save, show)


def plot_tissue_dotplot(df: pd.DataFrame,
                        cell_types: Optional[Sequence[str]] = None,
                        change: str = 'sign_dw',
                        n_tissues: int = 12,
                        min_genes: int = 1,
                        annotate: bool = True,
                        size_scale: float = 22.0,
                        figsize: Optional[Tuple[float, float]] = None,
                        ax: Optional[plt.Axes] = None,
                        title: Optional[str] = None,
                        save: Optional[str] = None,
                        show: bool = True):
    """Tissue of origin of the TRAs in a change class, as a tissue x cell type bubble plot.

    The stacked-bar version of this (``plot_tra_tissue_composition``) hides how few genes some
    populations contribute - TEC-iono has one down-regulated TRA, and a 100%-stacked bar makes that
    look like a strong tissue preference. Here dot **area** is the number of genes and dot
    **colour** is that tissue's share of the population's TRAs, so sparse populations read as
    sparse. Tissues are ordered by total gene count.

    Parameters
    ----------
    n_tissues
        Keep the top tissues by total count; the rest are pooled into 'other'.
    min_genes
        Do not draw dots below this count (0 is never drawn).
    size_scale
        Marker area per gene, in points^2.
    """

    order = _cell_type_order(df, cell_types)
    sub = df[(df['change'].astype(str) == change) & df['isTRA']].copy()
    sub = sub[sub['cell_type'].astype(str).isin(order)]
    if sub.empty:
        raise ValueError(f'no {change} TRAs in the selected cell types')

    top = sub['top1_tissue'].value_counts().head(n_tissues).index.tolist()
    sub['tissue'] = np.where(sub['top1_tissue'].isin(top), sub['top1_tissue'], 'other')
    tab = pd.crosstab(sub['tissue'], sub['cell_type'].astype(str)).reindex(columns=order,
                                                                          fill_value=0)
    row_order = [t for t in tab.sum(axis=1).sort_values(ascending=False).index if t != 'other']
    if 'other' in tab.index:
        row_order = row_order + ['other']
    tab = tab.reindex(row_order)
    frac = tab.div(tab.sum(axis=0).replace(0, np.nan), axis=1)      # share within a cell type

    if ax is None:
        if figsize is None:
            figsize = calc_figsize(width=30 + 12 * len(order), height=18 + 6 * len(tab), margin=0)
        fig, ax = plt.subplots(figsize=figsize)
    else:
        fig = ax.figure

    xs, ys, sizes, cols = [], [], [], []
    for j, ct in enumerate(order):
        for i, t in enumerate(tab.index):
            n = int(tab.loc[t, ct])
            if n < max(min_genes, 1):
                continue
            xs.append(j)
            ys.append(i)
            sizes.append(n * size_scale)
            cols.append(frac.loc[t, ct])

    cmap = sns.blend_palette(['white', thyAgeing_colors['teal']], as_cmap=True)
    sc = ax.scatter(xs, ys, s=sizes, c=cols, cmap=cmap, vmin=0, vmax=1,
                    edgecolor='k', linewidth=0.4, zorder=2)
    if annotate:
        for x, y, s, c in zip(xs, ys, sizes, cols):
            ax.text(x, y, f"{int(round(s / size_scale))}", ha="center", va="center",
                    fontsize=4.5, zorder=3, color="k" if c < 0.6 else "white")

    ax.set_xticks(range(len(order)))
    ax.set_xticklabels(order, rotation=45, ha='right', fontsize=7)
    ax.set_yticks(range(len(tab)))
    ax.set_yticklabels(tab.index, fontsize=7)
    ax.set_xlim(-0.7, len(order) - 0.3)
    ax.set_ylim(len(tab) - 0.4, -0.6)                               # first tissue at the top
    ax.set_axisbelow(True)
    ax.grid(True, color=thyAgeing_greys['grey1'], lw=0.4, zorder=0)
    for side in ('top', 'right'):
        ax.spines[side].set_visible(False)
    if title:
        ax.set_title(title)

    cb = fig.colorbar(sc, ax=ax, fraction=0.03, pad=0.03)
    cb.set_label('share of the population\'s TRAs', fontsize=6)
    cb.ax.tick_params(labelsize=6)
    cb.outline.set_linewidth(0.4)

    # size legend: a few representative gene counts
    counts = sorted({int(v) for v in tab.to_numpy().ravel() if v > 0})
    marks = [counts[0], counts[len(counts) // 2], counts[-1]] if len(counts) > 2 else counts
    handles = [Line2D([], [], marker='o', ls='none', markerfacecolor='white',
                      markeredgecolor='k', markeredgewidth=0.4,
                      markersize=np.sqrt(m * size_scale), label=str(m)) for m in dict.fromkeys(marks)]
    spacing = max(1.0, np.sqrt(max(marks) * size_scale) / 6.0)
    ax.legend(handles=handles, frameon=False, fontsize=6, title="n genes", title_fontsize=6,
              loc="upper left", bbox_to_anchor=(1.20, 1.0), labelspacing=spacing, borderpad=0.6)
    return _finish(fig, save, show)


def plot_tra_lollipop(df: pd.DataFrame,
                      cell_types: Optional[Sequence[str]] = None,
                      change: str = 'sign_dw',
                      top_n: int = 20,
                      sort_ascending: Optional[bool] = None,
                      n_tissues: int = 10,
                      ncols: int = 2,
                      panel_width: float = 55.0,
                      row_height: float = 3.6,
                      xlabel: str = 'log$_2$ fold change (adult vs infant)',
                      save: Optional[str] = None,
                      show: bool = True):
    """The genes themselves: one panel per cell type, set members ranked by log2 fold change.

    Each gene is a point at its log2FC, coloured by its GTEx tissue of origin, so the reader sees
    which antigens change and how strongly. Pass ``change='sign_up'`` for the genes that rise with
    age. Panels show at most ``top_n`` genes (the largest absolute fold changes) and the panel title
    reports how many there are in total. Panel rows are given heights proportional to the genes they
    hold, so a subset with four genes does not get the same vertical space as one with twenty.

    ``sort_ascending`` orders the rows. The default puts the largest fold change at the bottom of
    each panel whichever direction is plotted - ascending for a down-regulated set, descending for
    an up-regulated one - so the two directions can be read side by side.
    """

    order = _cell_type_order(df, cell_types)
    sub = df[(df['change'].astype(str) == change) & df['isTRA']].copy()
    sub = sub[sub['cell_type'].astype(str).isin(order)]
    if sub.empty:
        raise ValueError(f'no {change} TRAs in the selected cell types')

    top = sub['top1_tissue'].value_counts().head(n_tissues).index.tolist()
    sub['tissue'] = np.where(sub['top1_tissue'].isin(top), sub['top1_tissue'], 'other')
    tissues = list(top) + (['other'] if (sub['tissue'] == 'other').any() else [])
    pal = tissue_palette(tissues)

    panels = [(ct, sub[sub['cell_type'].astype(str) == ct]) for ct in order]
    panels = [(ct, d) for ct, d in panels if len(d)]
    nrows = int(np.ceil(len(panels) / ncols))
    rows = [panels[r * ncols:(r + 1) * ncols] for r in range(nrows)]
    row_max = [max(min(len(d), top_n) for _, d in r) for r in rows]

    fig, axs = plt.subplots(
        nrows, ncols, squeeze=False, gridspec_kw={'height_ratios': row_max},
        figsize=calc_figsize(width=min(183, panel_width * ncols + 30),
                             height=min(247, sum(row_height * m + 13 for m in row_max)),
                             margin=0))

    for k, (ct, d) in enumerate(panels):
        ax = axs[k // ncols][k % ncols]
        n_total = len(d)
        d = d.reindex(d['log2FoldChange'].abs().sort_values(ascending=False).index).head(top_n)
        asc = (d['log2FoldChange'].median() < 0) if sort_ascending is None else sort_ascending
        d = d.sort_values('log2FoldChange', ascending=asc)
        y = np.arange(len(d))
        ax.hlines(y, 0, d['log2FoldChange'], color=thyAgeing_greys['grey3'], lw=0.5, zorder=1)
        ax.scatter(d['log2FoldChange'], y, s=10, zorder=2, edgecolor='k', linewidth=0.4,
                   c=[pal[t] for t in d['tissue']])
        ax.set_yticks(y)
        ax.set_yticklabels(d['gene_name'], fontsize=5.5, style='italic')
        ax.axvline(0, color='k', lw=0.5)
        ax.set_xlabel(xlabel)
        gene_word = 'gene' if n_total == 1 else 'genes'
        shown = f'top {len(d)} of {n_total}' if n_total > len(d) else f'{n_total} {gene_word}'
        ax.set_title(f'{ct} ({shown})', fontsize=7)
        ax.set_ylim(-0.8, max(len(d), 1) - 0.2)
        # pad the non-zero side so the largest marker is not clipped by the axis
        lo = min(0.0, float(d['log2FoldChange'].min()))
        hi = max(0.0, float(d['log2FoldChange'].max()))
        pad = 0.08 * max(abs(lo), abs(hi), 1e-9)
        ax.set_xlim(lo - (pad if lo < 0 else 0.0), hi + (pad if hi > 0 else 0.0))
        for side in ('top', 'right'):
            ax.spines[side].set_visible(False)

    for k in range(len(panels), nrows * ncols):
        axs[k // ncols][k % ncols].axis('off')

    handles = [Patch(facecolor=pal[t], edgecolor='k', lw=0.4, label=t) for t in tissues]
    fig.legend(handles=handles, frameon=False, fontsize=6, title='GTEx tissue', title_fontsize=6,
               loc='upper left', bbox_to_anchor=(0.99, 0.95), handlelength=1.0)
    fig.tight_layout()
    return _finish(fig, save, show)


# the function used to be down-only; keep the old name as an alias so existing notebooks still run
plot_down_tra_lollipop = plot_tra_lollipop
